#!/usr/bin/env python3
"""
Nombres de lugar en los crudos históricos que solo traían códigos
(Centro de Descargas · catálogo histórico, oct-2026).

Agrega DES_DDE (departamento) tras COD_DDE, DES_MME (municipio) tras COD_MME y,
cuando hay con qué, DES_PP (puesto) tras COD_PP — la misma convención de
build.py. Rehace también la versión por puesto y, donde el catálogo ofrece
Excel, los .xlsx (votos como número; hoja nueva cada 1.048.575 filas).

⚠️ EL NOMBRE DEL PUESTO SALE DEL CENSO DE ESE AÑO, NO DEL DE 2026. El código de
puesto no es el mismo lugar de una elección a otra (Villavicencio conservaba un
tercio de sus códigos entre 2019 y 2026), así que nombrarlo con el georef
actual pondría votos en otro sitio sin que nada falle. Fuentes:
  2019 TER y JAL      censo 2019 (Registraduría, autoridades territoriales)
  2021 CLMJ (dic.)    Divipol 23-sep-2021
  2022 CONSU (mar.)   censo 2022
  2023 JAL            censo 2023
  2025 ATIP Magdalena censo 2025 de esa misma elección
  2025 CLMJ (oct.)    censo 2026 (el más cercano: corte de enero de 2026)
  2011 TER, 2015 TER y 2016 PLEB: NO llevan DES_PP. No existe censo por puesto
  de esos años y no hay con qué verificar que el código sea el mismo lugar.
Lo que el censo del año no nombra lleva el rótulo estructural de build.py
(entre corchetes, derivado del código: nunca se confunde con un nombre real).

Uso:
  python3 tools/build-csv-names/build_historicos.py              # los 10
  python3 tools/build-csv-names/build_historicos.py GCS_2016PLEB # uno
Entradas: el CSV de mesa de «FINAL SUBIDA GCS» y el _PUESTO.csv en
  Bases de datos/output_csv_names/historicos/orig/ (bajado de S3).
Salida: Bases de datos/output_csv_names/historicos/
"""
import csv
import json
import sys
import time
from pathlib import Path

ROOT = Path("/Users/ricardoruiz/ricardoruiz.co")
sys.path.insert(0, str(ROOT / "tools/build-csv-names"))
from build import load_maps, MUN_HISTORICOS, rotulo_puesto  # noqa: E402

SRC = ROOT / "Bases de datos/FINAL SUBIDA GCS"
OUT = ROOT / "Bases de datos/output_csv_names/historicos"
ORIG = OUT / "orig"
CENSOS = ROOT / "Bases de datos/censos-registraduria"
DIVIPOL21 = ROOT / "Bases de datos/Divipol 23.09.2021.xlsx"
MAX_FILAS = 1_048_575

# archivo → (fuente de nombres de puesto | None, ¿se ofrece Excel?)
ARCHIVOS = {
    "GCS_2011TER":      (None, False),
    "GCS_2015TER":      (None, False),
    "GCS_2019TER":      ("censo:2019", False),
    "GCS_2019JAL":      ("censo:2019", True),
    "GCS_2021CLMJ":     ("divipol2021", True),
    "GCS_2022CONSU":    ("censo:2022", True),
    "GCS_2023JAL":      ("censo:2023", True),
    "GCS_2025ATIP_MAG": ("censo:2025:GOBERNADOR MAGDALENA", True),
    "GCS_2025CLMJ":     ("censo:2026", True),
    "GCS_2016PLEB":     (None, True),
}


def arregla(c):
    """Texto UTF-8 que se leyó como Latin-1 en origen («POLÃ\x8dTICO» → «POLÍTICO»).
    Solo pasa en la atípica del Magdalena; si la reparación no da UTF-8 válido, se deja igual."""
    if "Ã" not in c:
        return c
    try:
        return c.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return c


def puestos_de(fuente):
    if fuente is None:
        return None
    nombres = {}
    if fuente == "divipol2021":
        import openpyxl
        wb = openpyxl.load_workbook(DIVIPOL21, read_only=True)
        for i, r in enumerate(wb.active.iter_rows(values_only=True), 1):
            if i <= 5 or r[0] is None or r[6] is None:
                continue
            nombres[str(r[0]).zfill(2) + str(r[1]).zfill(3) + str(r[2]).zfill(2) + str(r[3]).zfill(2)] = str(r[6]).strip()
        wb.close()
        return nombres
    _, anio, *tipo = fuente.split(":")
    d = json.loads((CENSOS / f"{anio}_datos_censo-electoral-min.json").read_text(encoding="utf-8"))
    rec = next(v for v in d.values() if isinstance(v, list)) if isinstance(d, dict) else d
    for r in rec:
        if tipo and r.get("tipo_eleccion") != tipo[0]:
            continue
        k = r["cod_depto"].zfill(2) + r["cod_mun"].zfill(3) + r["zona"].zfill(2) + r["cod_puesto"].zfill(2)
        nombres.setdefault(k, (r.get("puesto_votacion") or "").strip())
    return nombres


def enriquecer(entrada, salida, dep_n, mun_n, pue_n, guardar=False):
    """CSV → CSV con nombres. Devuelve (filas, votos, stats, encabezado, filas_para_excel)."""
    st = {"pp": 0, "rot": 0, "mun_rot": 0}
    filas_xlsx = [] if guardar else None
    with entrada.open(encoding="utf-8-sig", newline="") as fi, salida.open("w", encoding="utf-8-sig", newline="") as fo:
        rd = csv.reader(fi, delimiter=";")
        wr = csv.writer(fo, delimiter=";")
        h = next(rd)
        i_d, i_m, i_z, i_p, i_v = (h.index(c) for c in ("COD_DDE", "COD_MME", "COD_ZZ", "COD_PP", "NUM_VOT"))
        con_mun = "DES_MM" not in h          # 2011 ya trae el nombre del municipio en DES_MM
        # posiciones de inserción en orden descendente para no correr los índices
        ins = [(i_d, "DES_DDE")]
        if con_mun:
            ins.append((i_m, "DES_MME"))
        if pue_n is not None:
            ins.append((i_p, "DES_PP"))
        ins.sort(reverse=True)
        cab = list(h)
        for pos, nom in ins:
            cab.insert(pos + 1, nom)
        wr.writerow(cab)
        idx_v = cab.index("NUM_VOT")
        n = votos = 0
        for row in rd:
            if len(row) < len(h):
                row = row + [""] * (len(h) - len(row))
            dd = row[i_d].strip().zfill(2)
            mmm = row[i_m].strip().zfill(3)
            zz = row[i_z].strip().zfill(2)
            pp = row[i_p].strip().zfill(2)
            dn = dep_n.get(dd, "")
            mn = mun_n.get(f"{dd}-{mmm}", "") or MUN_HISTORICOS.get(f"{dd}-{mmm}", "")
            if not mn:
                mn = "[PAIS NO IDENTIFICADO]" if dd == "88" else "[MUNICIPIO NO IDENTIFICADO]"
                st["mun_rot"] += 1
            valores = {"DES_DDE": dn, "DES_MME": mn}
            if pue_n is not None:
                pn = pue_n.get(dd + mmm + zz + pp, "")
                if pn:
                    st["pp"] += 1
                else:
                    pn = rotulo_puesto(dd, mmm, zz, pp, dn, mn)
                    st["rot"] += 1
                valores["DES_PP"] = pn
            out = [arregla(c) for c in row] if "Ã" in ";".join(row) else list(row)
            for pos, nom in ins:
                out.insert(pos + 1, valores[nom])
            wr.writerow(out)
            n += 1
            try:
                votos += int(out[idx_v] or 0)
            except ValueError:
                pass
            if filas_xlsx is not None:
                filas_xlsx.append(out)
                if len(filas_xlsx) > 4_000_000:   # no cabe en memoria ni en Excel: no se pide
                    filas_xlsx = None
            if n % 2_000_000 == 0:
                print(f"     {n:,} filas…", flush=True)
    return n, votos, st, cab, filas_xlsx


def escribir_xlsx(ruta, cab, filas):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    wb = Workbook(write_only=True)
    i_v = cab.index("NUM_VOT")
    negrita = Font(bold=True)
    hoja, en_hoja, k = None, 0, 0
    for f in filas:
        if hoja is None or en_hoja >= MAX_FILAS:
            k += 1
            hoja = wb.create_sheet("Datos" if k == 1 else f"Datos {k}")
            hoja.freeze_panes = "A2"
            enc = []
            for c in cab:
                x = WriteOnlyCell(hoja, value=c)
                x.font = negrita
                enc.append(x)
            hoja.append(enc)
            en_hoja = 0
        g = list(f)
        try:
            g[i_v] = int(g[i_v])
        except (ValueError, TypeError):
            pass
        hoja.append(g)
        en_hoja += 1
    wb.save(ruta)
    return k


def votos_de(ruta):
    with ruta.open(encoding="utf-8-sig", newline="") as f:
        rd = csv.reader(f, delimiter=";")
        h = next(rd)
        i = h.index("NUM_VOT")
        n = v = 0
        for r in rd:
            n += 1
            try:
                v += int(r[i] or 0)
            except (ValueError, IndexError):
                pass
    return n, v


def main():
    pedidos = [a for a in sys.argv[1:] if not a.startswith("-")] or list(ARCHIVOS)
    dep_n, mun_n, _ = load_maps()
    OUT.mkdir(parents=True, exist_ok=True)
    for base in pedidos:
        fuente, con_xlsx = ARCHIVOS[base]
        t0 = time.time()
        pue_n = puestos_de(fuente)
        print(f"\n{base} · puestos: {fuente or 'sin nombre (no hay censo de ese año)'}", flush=True)
        for nivel, entrada in (("mesa", SRC / f"{base}.csv"), ("puesto", ORIG / f"{base}_PUESTO.csv")):
            sufijo = "" if nivel == "mesa" else "_PUESTO"
            salida = OUT / f"{base}{sufijo}.csv"
            n, v, st, cab, filas = enriquecer(entrada, salida, dep_n, mun_n, pue_n, con_xlsx)
            n0, v0 = votos_de(entrada)
            ok = "✓" if (n, v) == (n0, v0) else "✗ NO CUADRA"
            linea = f"   {nivel}: {n:,} filas · {v:,} votos {ok} (original {n0:,} · {v0:,})"
            if pue_n is not None:
                linea += f" · puesto con nombre {st['pp'] / n * 100:.1f} %"
            if st["mun_rot"]:
                linea += f" · {st['mun_rot']:,} filas sin municipio identificado"
            print(linea, flush=True)
            if (n, v) != (n0, v0):
                sys.exit(f"{base} {nivel}: el enriquecido no cuadra con el original")
            if con_xlsx:
                if filas is None:
                    sys.exit(f"{base} {nivel}: demasiadas filas para Excel")
                hojas = escribir_xlsx(OUT / f"{base}{sufijo}.xlsx", cab, filas)
                print(f"   {nivel}: Excel en {hojas} hoja(s)", flush=True)
        print(f"   {time.time() - t0:.0f} s", flush=True)


if __name__ == "__main__":
    main()
