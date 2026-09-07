#!/usr/bin/env python3
"""
Borrador del tema 7 (Inteligencia artificial y tecnologías emergentes) para el
"Diagnóstico del sector TIC en Colombia — insumos para el PND 2026-2030".

A diferencia de la nota de insumos (build_nota_temas.py), esto NO es un informe
propio: es la sección lista para incorporarse al documento que arma la
coordinación. Por eso no lleva portada, control documental ni tabla de
contenido, y el formato replica el del borrador de septiembre de 2026 en vez de
las pautas de informe técnico de la DENDD:

    Arial Narrow 10, justificado, interlineado 257 automático (~1,07),
    encabezados con los estilos Heading1/Heading2 del propio borrador.

Se respeta la estructura de cinco preguntas fijada en la metodología (situación
actual, evolución, brechas, retos 2026-2030 y mensaje para el PND), el límite de
dos páginas y la indicación de priorizar entre tres y cinco indicadores clave.

Las cifras de nivel A provienen del análisis de este contrato sobre el Plan de
Acción y Seguimiento y los reportes de SisCONPES del CONPES 4144 (corte del 21
de julio de 2026) y sobre los datos abiertos del sector. Las de nivel B llevan
su fuente publicada y deben cotejarse contra el documento antes de la versión
final, según el protocolo de la nota de insumos.

Genera dos archivos: el borrador de la sección, que es el que se comparte, y
una nota de verificación aparte para uso interno de la coordinación.

Uso:
    python3 build_tema7.py [--salida RUTA.docx] [--salida-nota RUTA.docx] [--pdf]
"""

import argparse
import os
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FUENTE = "Arial Narrow"
NEGRO = RGBColor(0, 0, 0)
INTERLINEADO = "257"  # el mismo w:spacing/@w:line del borrador

AQUI = os.path.dirname(os.path.abspath(__file__))

# Dos archivos separados a propósito: el primero es el que se envía por correo
# al equipo y solo contiene la sección; el segundo es material de trabajo y no
# debe viajar dentro del documento que se comparte.
SALIDA_DEFECTO = os.path.join(AQUI, "Tema7-IA-y-tecnologias-emergentes-borrador-v1.docx")
SALIDA_NOTA = os.path.join(AQUI, "Tema7-nota-de-verificacion-interna.docx")

SOFFICE = next((r for r in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
                            "/usr/bin/soffice", "/usr/bin/libreoffice")
                if os.path.exists(r)), None)


# --------------------------------------------------------------------------
# Formato del documento de destino
# --------------------------------------------------------------------------

def _interlineado(p):
    """Replica el <w:spacing w:line="257" w:lineRule="auto"/> del borrador.
    python-docx no expone lineRule="auto" con un valor crudo, así que se
    escribe directamente sobre el pPr."""
    pPr = p._p.get_or_add_pPr()
    for previo in pPr.findall(qn("w:spacing")):
        pPr.remove(previo)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:line"), INTERLINEADO)
    sp.set(qn("w:lineRule"), "auto")
    sp.set(qn("w:after"), "120")
    pPr.append(sp)


def configurar(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(10)
    normal.font.color.rgb = NEGRO
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    for nivel, tam in ((1, 13), (2, 11)):
        st = doc.styles[f"Heading {nivel}"]
        st.font.name = FUENTE
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = NEGRO
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    for lado, cm in (("top", 2.54), ("bottom", 2.54), ("left", 2.54), ("right", 2.54)):
        setattr(doc.sections[0], f"{lado}_margin", Cm(cm))


def p(doc, texto, negrita=False, tam=10):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.font.name = FUENTE
    run.font.size = Pt(tam)
    run.font.bold = negrita
    run.font.color.rgb = NEGRO
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _interlineado(par)
    return par


def subtitulo(doc, texto):
    """Los rótulos de las cinco preguntas van en negrita sobre su propia línea,
    como en el tema 2 del borrador, no como encabezado numerado."""
    par = p(doc, texto, negrita=True)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.keep_with_next = True
    return par


def titulo(doc, texto, nivel=2):
    h = doc.add_heading(texto, level=nivel)
    for run in h.runs:
        run.font.name = FUENTE
        run.font.color.rgb = NEGRO
    return h


def vineta(doc, texto):
    par = doc.add_paragraph(style="List Bullet")
    run = par.add_run(texto)
    run.font.name = FUENTE
    run.font.size = Pt(10)
    run.font.color.rgb = NEGRO
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _interlineado(par)
    return par


def tabla(doc, titulo_tabla, encabezados, filas, fuente_nota, anchos):
    cap = doc.add_paragraph()
    r = cap.add_run(titulo_tabla)
    r.font.name = FUENTE
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = NEGRO
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True

    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)

    def celda(c, texto, negrita=False):
        c.text = ""
        par = c.paragraphs[0]
        run = par.add_run(texto)
        run.font.name = FUENTE
        run.font.size = Pt(8.5)
        run.font.bold = negrita
        run.font.color.rgb = NEGRO
        par.paragraph_format.space_after = Pt(0)

    for i, texto in enumerate(encabezados):
        celda(t.rows[0].cells[i], texto, negrita=True)
    for fila in filas:
        celdas = t.add_row().cells
        for i, texto in enumerate(fila):
            celda(celdas[i], texto)

    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, ancho in zip(grid.findall(qn("w:gridCol")), anchos):
            col.set(qn("w:w"), str(int(Cm(ancho).twips)))
    for fila in t.rows:
        for i, ancho in enumerate(anchos):
            fila.cells[i].width = Cm(ancho)
            tcW = fila.cells[i]._tc.get_or_add_tcPr().get_or_add_tcW()
            tcW.set(qn("w:w"), str(int(Cm(ancho).twips)))
            tcW.set(qn("w:type"), "dxa")

    nota = doc.add_paragraph()
    r = nota.add_run(f"Fuente: {fuente_nota}")
    r.font.name = FUENTE
    r.font.size = Pt(8)
    r.font.color.rgb = NEGRO
    nota.paragraph_format.space_after = Pt(8)
    return t


# --------------------------------------------------------------------------
# Contenido de la sección
# --------------------------------------------------------------------------

def seccion(doc):
    titulo(doc, "7. Inteligencia artificial y tecnologías emergentes", 2)

    subtitulo(doc, "Situación actual")
    p(doc,
      "Colombia cuenta desde febrero de 2025 con una política nacional de "
      "inteligencia artificial. El documento CONPES 4144 organiza 106 acciones "
      "alrededor de seis objetivos y cuenta con financiación completa para su "
      "horizonte. Su ejecución, sin embargo, avanza por debajo de lo previsto: "
      "al corte del 21 de julio de 2026 el avance físico promedio de las "
      "acciones era del 43,3 %, con una disciplina de reporte del 87,7 %. El "
      "freno, por tanto, no es presupuestal sino de implementación. "
      "Veinticuatro acciones se encuentran finalizadas —entre ellas la "
      "producción de la estadística oficial de pobreza digital, que hoy sirve "
      "de referente de resultado al propio sector—, pero quince de las "
      "dieciocho acciones que vencían en 2025 cerraron el año sin cumplimiento "
      "total.")
    p(doc,
      "El contraste relevante para el diagnóstico es que la adopción social de "
      "la tecnología avanza más rápido que la política llamada a gobernarla. "
      "Según la Encuesta de Tecnologías de la Información y las Comunicaciones "
      "en Hogares 2024, el 18,0 % de las personas de cinco años o más que "
      "usaron internet lo hicieron para utilizar herramientas de inteligencia "
      "artificial. En la comparación regional, el Índice Latinoamericano de "
      "Inteligencia Artificial 2025 ubica a Colombia en el grupo de países "
      "«adoptantes», por debajo del umbral que separa a los líderes de la "
      "región.")

    tabla(doc, "Tabla 1. Indicadores clave del tema",
          ["Indicador", "Valor", "Corte", "Fuente"],
          [["Avance físico promedio de las acciones del CONPES 4144",
            "43,3 %", "21 jul. 2026", "SisCONPES"],
           ["Acciones vencidas en 2025 sin cumplimiento total",
            "15 de 18", "21 jul. 2026", "SisCONPES"],
           ["Indicadores de resultado en el Plan de Acción y Seguimiento",
            "0 de 106", "2025", "PAS CONPES 4144"],
           ["Entidades que concentran las acciones de la política",
            "3 entidades en el 80,2 %", "2025", "PAS CONPES 4144"],
           ["Personas usuarias de internet que usaron herramientas de IA",
            "18,0 %", "2024", "DANE, ENTIC Hogares"]],
          "elaboración propia con base en el Plan de Acción y Seguimiento y los "
          "reportes de SisCONPES del CONPES 4144 facilitados por la Dirección, y "
          "en el boletín técnico de la ENTIC Hogares 2024 del DANE.",
          anchos=[7.2, 3.4, 2.2, 3.2])

    subtitulo(doc, "Evolución")
    p(doc,
      "Entre los dos últimos gobiernos el cambio principal es de estatus del "
      "tema. En 2019 la inteligencia artificial era un componente dentro de una "
      "política más amplia de transformación digital (CONPES 3975); en 2025 "
      "pasó a tener política propia, con arquitectura de gobernanza, enfoque "
      "basado en riesgos e institucionalidad específica. La comparación "
      "sistemática de ambos documentos muestra además que la formulación de "
      "2025 incorporó de manera explícita los enfoques territorial y "
      "diferencial que su antecesora prácticamente no contenía, lo que "
      "constituye un avance real de diseño. A ello se suma la aparición de "
      "instrumentos institucionales que en 2019 no existían, como el "
      "Observatorio Nacional de Inteligencia Artificial y la red nacional de "
      "datos e inteligencia artificial. El salto de los últimos años, en suma, "
      "es de formulación; la ejecución todavía no lo acompaña.")

    subtitulo(doc, "Brechas")
    p(doc,
      "Las brechas del tema son de tres órdenes y ninguna se resuelve con más "
      "presupuesto.")
    vineta(doc,
           "De medición. El Plan de Acción y Seguimiento contiene 76 "
           "indicadores de gestión y 30 de producto, y ninguno de resultado. La "
           "política puede documentar cuánta actividad realizó, pero no si esa "
           "actividad produjo algún efecto sobre la productividad, el empleo o "
           "el acceso. Es la brecha más determinante, porque condiciona "
           "cualquier evaluación futura.")
    vineta(doc,
           "Institucional. Cincuenta y seis entidades participan en la "
           "política, pero tres concentran el 80,2 % de las acciones y la mitad "
           "figura en una sola. La capacidad efectiva de ejecución está "
           "concentrada y la corresponsabilidad es, en buena parte, nominal. En "
           "el mismo sentido, la asignación de recursos privilegia la "
           "infraestructura sobre las capacidades en una proporción cercana a "
           "quince a uno, sin que el documento explicite la hipótesis que "
           "conecta una cosa con la otra.")
    vineta(doc,
           "Distributiva. Las acciones asociadas al enfoque de derechos y a los "
           "grupos poblacionales priorizados por la propia política figuran "
           "entre las de menor avance. El hallazgo se replica en el terreno: en "
           "el registro de participantes certificados en formación en "
           "inteligencia artificial, la brecha de género se produce en el "
           "acceso al programa y no en la permanencia dentro de él, que son dos "
           "problemas de política distintos y exigen respuestas distintas.")

    subtitulo(doc, "Retos 2026-2030")
    p(doc,
      "El primer reto es de calendario y es inmediato: cuarenta y ocho acciones "
      "de la política terminan su vigencia en 2026, de modo que el próximo Plan "
      "recibe el instrumento en su momento de mayor exigencia y con la "
      "necesidad de decidir qué se prorroga, qué se cierra y qué se rediseña. "
      "Asociado a él, la estrategia de gobernanza anticipatoria está vencida "
      "con un avance del 38 %, y su cierre puede apalancarse en la "
      "consolidación del Observatorio Nacional de Inteligencia Artificial "
      "durante el período.")
    p(doc,
      "El segundo es estructural y excede a esta política: el desfase entre la "
      "velocidad del cambio tecnológico y el ciclo de la política pública. Los "
      "instrumentos se formulan con horizontes de cinco años sobre tecnologías "
      "que se transforman en meses, lo que exige incorporar mecanismos "
      "explícitos de actualización y revisión periódica en lugar de esperar al "
      "siguiente documento de política. El tercero es de autonomía: la "
      "capacidad de cómputo, la disponibilidad de datos de entrenamiento "
      "pertinentes al contexto nacional y la retención de talento especializado "
      "determinarán si el país desarrolla capacidad propia o consolida una "
      "posición de dependencia tecnológica.")

    subtitulo(doc, "Mensaje para el PND")
    p(doc,
      "El problema central de este tema no es la falta de recursos ni la "
      "calidad de la formulación, que mejoró de manera verificable frente al "
      "instrumento anterior. El problema es que la política nacional de "
      "inteligencia artificial no está en condiciones de demostrar resultados: "
      "avanza al 43,3 % con financiación completa, y su sistema de seguimiento "
      "no contiene un solo indicador de resultado. Si el Plan Nacional de "
      "Desarrollo 2026-2030 hereda ese diseño de medición, hereda también la "
      "imposibilidad de establecer si la inversión pública en inteligencia "
      "artificial modificó algo.")
    p(doc,
      "En consecuencia, el Plan debería incorporar tres definiciones. Primera, "
      "dotar a la política de indicadores de resultado con línea base y fuente "
      "identificada, condición sin la cual ninguna evaluación posterior es "
      "posible. Segunda, corregir la concentración institucional de la "
      "ejecución y hacer efectiva la corresponsabilidad de las entidades que "
      "hoy figuran de manera nominal. Tercera, proteger de manera explícita las "
      "acciones de equidad y enfoque diferencial, hoy las más rezagadas, "
      "mientras el rezago todavía es corregible. La oportunidad de fondo es "
      "cerrar la distancia entre una ciudadanía que ya incorporó la "
      "inteligencia artificial a su uso cotidiano de internet y un Estado que "
      "apenas comienza a gobernarla.")


def fuentes_seccion(doc):
    subtitulo(doc, "Fuentes de esta sección")
    for texto in (
        "Comisión Económica para América Latina y el Caribe y Centro Nacional de "
        "Inteligencia Artificial. (2025). Índice Latinoamericano de Inteligencia "
        "Artificial (ILIA) 2025. https://repositorio.cepal.org/handle/11362/82514",
        "Departamento Administrativo Nacional de Estadística. (2025). Encuesta de "
        "Tecnologías de la Información y las Comunicaciones en Hogares 2024: boletín "
        "técnico. https://www.dane.gov.co/files/operaciones/ENTIC/bol-ENTICHogares-2024.pdf",
        "Departamento Nacional de Planeación. (2019). Documento CONPES 3975. Política "
        "nacional para la transformación digital e inteligencia artificial. DNP.",
        "Departamento Nacional de Planeación. (2025). Documento CONPES 4144. Política "
        "nacional de inteligencia artificial. DNP. "
        "https://colaboracion.dnp.gov.co/CDT/Conpes/Económicos/4144.pdf",
        "Departamento Nacional de Planeación. (2025). Anexo A. Plan de Acción y "
        "Seguimiento del documento CONPES 4144 [instrumento de seguimiento, SisCONPES]. DNP.",
        "Departamento Nacional de Planeación. (2026). Reporte y revisión del documento "
        "CONPES 4144 [módulo de SisCONPES, corte del 21 de julio de 2026]. DNP.",
        "Ministerio de Tecnologías de la Información y las Comunicaciones. (2026). "
        "Participantes certificados en inteligencia artificial – SENATIC [conjunto de "
        "datos, identificador m2uu-cu4q]. Portal de Datos Abiertos de Colombia. "
        "https://www.datos.gov.co/resource/m2uu-cu4q.json",
    ):
        par = doc.add_paragraph()
        r = par.add_run(texto)
        r.font.name = FUENTE
        r.font.size = Pt(8.5)
        r.font.color.rgb = NEGRO
        pf = par.paragraph_format
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.9)
        pf.space_after = Pt(3)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def nota_coordinacion(doc):
    """Documento de trabajo aparte: acompaña al borrador pero no se incorpora
    al diagnóstico ni viaja dentro del archivo que se comparte por correo."""
    titulo(doc, "Tema 7 · Nota de verificación para la coordinación", 2)
    p(doc,
      "Documento de trabajo que acompaña al borrador de la sección «7. "
      "Inteligencia artificial y tecnologías emergentes». No hace parte del "
      "texto del diagnóstico.")
    p(doc,
      "Estado de verificación de las cifras empleadas, según el protocolo de la "
      "nota de insumos de este contrato.")
    p(doc,
      "Nivel A, reproducibles desde la fuente primaria por este contrato: el "
      "avance físico del 43,3 %, la disciplina de reporte del 87,7 %, las 24 "
      "acciones finalizadas, las 15 de 18 acciones vencidas sin cumplimiento, "
      "la composición del sistema de indicadores (76 de gestión, 30 de producto "
      "y ninguno de resultado), la concentración del 80,2 % en tres entidades, "
      "las 48 acciones que vencen en 2026, el 38 % de la estrategia "
      "anticipatoria, la proporción de quince a uno entre infraestructura y "
      "capacidades, y la brecha de género por acceso y no por permanencia en el "
      "registro de certificación. Todas provienen del Plan de Acción y "
      "Seguimiento y de los reportes de SisCONPES facilitados por la Dirección, "
      "y del procesamiento de datos abiertos documentado en el Anexo A de la "
      "Entrega 3.")
    p(doc,
      "Nivel B, publicadas y con enlace identificado pero pendientes de cotejo "
      "directo contra el documento: el 18,0 % de uso de herramientas de IA de "
      "la ENTIC Hogares 2024 y la clasificación de Colombia en el grupo de "
      "«adoptantes» del ILIA 2025. Antes de la versión final conviene abrir "
      "ambos documentos y confirmar el dato y su año de referencia; son dos "
      "consultas de pocos minutos.")
    p(doc,
      "Dos precisiones pendientes que dependen de la Dirección. La primera es "
      "el estado real de operación del Observatorio Nacional de Inteligencia "
      "Artificial, que la sección menciona en «Evolución» y en «Retos» sin "
      "cifra porque no se identificó fuente publicada. La segunda es si el "
      "reporte de SisCONPES con corte al cierre de 2026 estará disponible antes "
      "de la versión final del diagnóstico; de estarlo, conviene recalcular el "
      "avance sobre reportes aprobados y actualizar la Tabla 1, dejando "
      "constancia del cambio de corte.")
    p(doc,
      "Sobre extensión: el cuerpo de la sección ocupa dos páginas en el formato "
      "del borrador (Arial Narrow 10, justificado). Si la coordinación necesita "
      "recortar, el párrafo de mayor holgura es el segundo de «Retos "
      "2026-2030», cuyos dos últimos retos pueden comprimirse en una frase sin "
      "perder el argumento.")


# --------------------------------------------------------------------------

def construir(salida, salida_nota):
    doc = Document()
    configurar(doc)
    seccion(doc)
    fuentes_seccion(doc)
    doc.core_properties.title = "Tema 7. Inteligencia artificial y tecnologías emergentes"
    doc.core_properties.subject = "Diagnóstico del sector TIC — insumos PND 2026-2030"
    doc.save(salida)

    nota = Document()
    configurar(nota)
    nota_coordinacion(nota)
    nota.core_properties.title = "Tema 7. Nota de verificación para la coordinación"
    nota.core_properties.subject = "Diagnóstico del sector TIC — insumos PND 2026-2030"
    nota.save(salida_nota)
    return salida, salida_nota


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--salida", default=SALIDA_DEFECTO)
    ap.add_argument("--salida-nota", default=SALIDA_NOTA)
    ap.add_argument("--pdf", action="store_true")
    args = ap.parse_args()

    rutas = construir(args.salida, args.salida_nota)
    for ruta in rutas:
        print(ruta)
    if args.pdf:
        if not SOFFICE:
            print("aviso: no se encontró LibreOffice; se omite el PDF")
            return 0
        for ruta in rutas:
            subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                            "--outdir", os.path.dirname(ruta), ruta],
                           check=True, capture_output=True)
            print(ruta.replace(".docx", ".pdf"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
