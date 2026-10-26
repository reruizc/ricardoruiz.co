#!/usr/bin/env python3
"""
Datos del territorio para el panel 12 de Candidato 360 («Su territorio, mes a
mes»): empleo del DANE por ciudad y educación del MEN por municipio, en un JSON
chico por código DANE. La seguridad NO pasa por aquí: el panel lee directo los
JSON del hub de Policía (congreso-2026/output/ponal/), que ya están publicados.

Empleo · los dos anexos que el DANE publica cada mes:

  anex-GEIH-<mes><año>.xlsx
    · «Total 23 ciudades A.M. Trim»  → trimestre móvil: TD, TO, TGP (23 ciudades)
    · «Año_móvil_32_ciudades»        → año móvil: TD, TO, TGP (las 32 capitales)
    · «Total nacional Trim»          → la referencia nacional
  anex-GEIHEISS-<trim><año>.xlsx
    · «Ciudades»                     → informalidad, trimestre móvil (23 ciudades)

Cada ciudad lleva los municipios de su área metropolitana tal como la define
la GEIH (Medellín = Valle de Aburrá; Barranquilla = con Soledad; Cúcuta = con
Villa del Rosario, Los Patios y El Zulia…). Un concejo de Bello lee el dato de
«Medellín A.M.» y la página lo dice; un municipio que no está en ninguna lee el
de la capital de su departamento, rotulado como referencia.

Educación · `Bases de datos/indicadores-mun/indicadores-mun.json` (Sprint E del
Lab: cobertura neta y deserción del MEN, 2018-2024). ⚠ El MEN reporta por
entidad territorial certificada: un municipio no certificado trae el dato de su
departamento y va marcado `aprox`. Bogotá venía marcado como aproximado y no lo
es (es su propia ETC): se corrige aquí.

    python3 tools/candidato-360/territorio/build.py \
        "Bases de datos/dane-empleo/raw/anex-GEIH-ago2026.xlsx" \
        "Bases de datos/dane-empleo/raw/anex-GEIHEISS-jun-ago2026.xlsx"

Salida: Bases de datos/dane-empleo/territorio.json
S3:     congreso-2026/output/candidato-360/territorio.json (gzip)
"""
import json, os, re, sys, unicodedata
import openpyxl

ROOT = "/Users/ricardoruiz/ricardoruiz.co"
OUT = os.path.join(ROOT, "Bases de datos", "dane-empleo", "territorio.json")
EDU = os.path.join(ROOT, "Bases de datos", "indicadores-mun", "indicadores-mun.json")
N_PER = 26   # cuántos periodos se guardan (dos años y un poco)

CAPITALES = {
    "BOGOTA": "11001", "MEDELLIN": "05001", "CALI": "76001", "BARRANQUILLA": "08001", "BUCARAMANGA": "68001",
    "MANIZALES": "17001", "PASTO": "52001", "PEREIRA": "66001", "CUCUTA": "54001", "IBAGUE": "73001",
    "MONTERIA": "23001", "CARTAGENA": "13001", "VILLAVICENCIO": "50001", "TUNJA": "15001", "FLORENCIA": "18001",
    "POPAYAN": "19001", "VALLEDUPAR": "20001", "QUIBDO": "27001", "NEIVA": "41001", "RIOHACHA": "44001",
    "SANTA MARTA": "47001", "ARMENIA": "63001", "SINCELEJO": "70001", "ARAUCA": "81001", "YOPAL": "85001",
    "MOCOA": "86001", "SAN ANDRES": "88001", "LETICIA": "91001", "INIRIDA": "94001",
    "SAN JOSE DEL GUAVIARE": "95001", "MITU": "97001", "PUERTO CARRENO": "99001",
}
# Áreas metropolitanas de la GEIH (municipios que entran en el dato «A.M.»).
AREA = {
    "05001": ["05001", "05079", "05088", "05129", "05212", "05266", "05308", "05360", "05380", "05631"],
    "76001": ["76001", "76892"],
    "08001": ["08001", "08758"],
    "68001": ["68001", "68276", "68307", "68547"],
    "17001": ["17001", "17873"],
    "66001": ["66001", "66170", "66400"],
    "54001": ["54001", "54261", "54405", "54874"],
}
MESES = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AGO": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DIC": 12}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z ]", " ", s)).strip()


def ciudad_de(nombre):
    n = norm(nombre).replace(" A M", "").replace(" D C", "").strip()
    return CAPITALES.get(n)


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def etiquetas_trim(fila_anio, fila_per):
    """Une la fila de años (solo trae el año en el primer periodo de cada uno)
    con la de periodos: «Jun - Ago 2026»."""
    out, anio = [], None
    for a, p in zip(fila_anio, fila_per):
        if a is not None and re.match(r"^\d{4}", str(a).strip()):
            anio = str(a).strip()[:4]
        out.append(f"{str(p).strip()} {anio}" if p is not None and anio else None)
    return out


def bloques(ws, tipo):
    """Recorre una hoja con un bloque por ciudad: nombre · … · «Concepto» ·
    fila de periodos · filas de tasas."""
    filas = list(ws.iter_rows(values_only=True))
    res, i = {}, 0
    while i < len(filas):
        c0 = filas[i][0]
        if c0 and str(c0).strip() == "Concepto":
            # el nombre de la ciudad está unas filas arriba
            nombre = None
            for k in range(i - 1, max(i - 5, -1), -1):
                v = filas[k][0]
                if v and not str(v).startswith(("Serie", "Año móvil")):
                    nombre = str(v).strip(); break
            per = filas[i + 1]
            labels = etiquetas_trim(filas[i], per) if tipo == "trim" else [str(p).strip() if p else None for p in per]
            datos = {}
            for k in range(i + 2, min(i + 14, len(filas))):
                lab = norm(filas[k][0])
                for clave, pat in (("tgp", "TASA GLOBAL"), ("to", "TASA DE OCUPACION"), ("td", "TASA DE DESOCUPACION")):
                    if lab.startswith(pat):
                        datos[clave] = [num(v) for v in filas[k]]
            idx = [j for j in range(1, len(labels)) if labels[j] and datos.get("td") and j < len(datos["td"]) and datos["td"][j] is not None]
            idx = idx[-N_PER:]
            res[nombre] = {"p": [labels[j] for j in idx], **{k: [v[j] for j in idx] for k, v in datos.items()}}
            i += 2
        i += 1
    return res


def informalidad(ws):
    filas = list(ws.iter_rows(values_only=True))
    hdr = next(k for k, f in enumerate(filas) if f[0] and str(f[0]).strip() == "Concepto")
    labels = etiquetas_trim(filas[hdr], filas[hdr + 1])
    res, ciudad, ocup = {}, None, None
    for f in filas[hdr + 2:]:
        if f[0]:
            ciudad = str(f[0]).strip()
        lab = norm(f[1])
        if lab == "POBLACION OCUPADA":
            ocup = [num(v) for v in f]
        elif lab == "INFORMAL" and ciudad and ocup:
            inf = [num(v) for v in f]
            idx = [j for j in range(2, len(labels)) if labels[j] and inf[j] and ocup[j]][-N_PER:]
            res[ciudad] = {"p": [labels[j] for j in idx], "inf": [round(100 * inf[j] / ocup[j], 1) for j in idx]}
    return res


def main(geih, eiss):
    wb = openpyxl.load_workbook(geih, read_only=True, data_only=True)
    trim = bloques(wb["Total 23 ciudades A.M. Trim"], "trim")
    anual = bloques(wb["Año_móvil_32_ciudades"], "anual")
    nac = bloques(wb["Total nacional Trim"], "trim")
    wbi = openpyxl.load_workbook(eiss, read_only=True, data_only=True)
    inf = informalidad(wbi["Ciudades"])

    ciudades = {}
    for nombre, d in list(trim.items()) + list(anual.items()):
        code = ciudad_de(nombre)
        if not code:
            continue
        c = ciudades.setdefault(code, {"n": nombre, "area": AREA.get(code, [code])})
        c["trim" if nombre in trim and d is trim[nombre] else "anual"] = d
    for nombre, d in inf.items():
        code = ciudad_de(nombre)
        if code in ciudades:
            ciudades[code]["inf"] = d
    nacional = next(iter(nac.values()))

    ind = json.load(open(EDU, encoding="utf-8"))
    educacion = {}
    for code, m in ind["muns"].items():
        dd = m.get("datos", {})
        meta = dd.get("_meta", {})
        fila = {}
        for k, corto in (("cobertura_neta", "cob"), ("desercion", "des")):
            serie = dd.get(k) or {}
            if serie:
                anios = sorted(serie)
                fila[corto] = [anios[0], [serie[a] for a in anios]]
        if not fila:
            continue
        if meta.get("cobertura_neta") and code != "11001":
            fila["aprox"] = 1
        educacion[code] = fila
    out = {
        "v": os.path.basename(geih).replace("anex-GEIH-", "").replace(".xlsx", ""),
        "fuente": "DANE · Gran Encuesta Integrada de Hogares (GEIH), anexos de mercado laboral y de informalidad",
        "archivos": [os.path.basename(geih), os.path.basename(eiss)],
        "nacional": nacional,
        "ciudades": ciudades,
        "educacion": educacion,
        "educacion_fuente": "Ministerio de Educación · estadísticas por entidad territorial certificada (datos.gov.co sras-4t5p), 2018-2024",
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    con_trim = sum(1 for c in ciudades.values() if "trim" in c)
    print(f"→ {OUT} · {os.path.getsize(OUT)//1024} KB · {len(ciudades)} ciudades ({con_trim} con trimestre móvil, {sum(1 for c in ciudades.values() if 'inf' in c)} con informalidad)")
    faltan = sorted(set(CAPITALES.values()) - set(ciudades))
    if faltan:
        print("⚠ capitales sin dato:", faltan)
    b = ciudades["11001"]
    print("control Bogotá:", b["trim"]["p"][-1], "TD", b["trim"]["td"][-1], "· año móvil", b["anual"]["p"][-1], b["anual"]["td"][-1], "· informalidad", b["inf"]["p"][-1], b["inf"]["inf"][-1])
    print("control nacional:", nacional["p"][-1], "TD", nacional["td"][-1])
    print(f"educación: {len(educacion)} municipios · {sum(1 for e in educacion.values() if e.get('aprox'))} con dato departamental · Bogotá", educacion["11001"])


if __name__ == "__main__":
    main(*sys.argv[1:3])
