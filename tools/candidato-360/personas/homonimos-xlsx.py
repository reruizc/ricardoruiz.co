# Los homónimos de un departamento, en un Excel para revisar a mano.
#
#   node tools/candidato-360/personas/similares.mjs        # deja salida/similares.json
#   python3 tools/candidato-360/personas/homonimos-xlsx.py 1 16
#
# Cada argumento es un código de departamento (1 = Antioquia, 16 = Bogotá). Deja
# un libro por departamento en salida/, con dos hojas: «Cómo llenarlo» y «Casos».
# La columna B es lo único que se llena: X (misma persona), N (homónimos) o ?.
import json, os, re, sys, zipfile, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, 'salida')
grupos = json.load(open(os.path.join(SALIDA, 'similares.json'), encoding='utf-8'))
DEPN = {
    '1': 'Antioquia', '3': 'Atlántico', '5': 'Bolívar', '7': 'Boyacá', '9': 'Caldas',
    '11': 'Cauca', '12': 'Cesar', '13': 'Córdoba', '15': 'Cundinamarca', '16': 'Bogotá D.C.',
    '17': 'Chocó', '19': 'Huila', '21': 'Magdalena', '23': 'Nariño', '24': 'Risaralda',
    '25': 'Norte de Santander', '26': 'Quindío', '27': 'Santander', '28': 'Sucre',
    '29': 'Tolima', '31': 'Valle del Cauca', '40': 'Arauca', '44': 'Caquetá',
    '46': 'Casanare', '48': 'La Guajira', '50': 'Guainía', '52': 'Meta', '54': 'Guaviare',
    '56': 'San Andrés', '60': 'Amazonas', '64': 'Putumayo', '68': 'Vaupés', '72': 'Vichada',
}

INK = '102238'; VERDE = '147C70'; CORAL = 'EF705D'; LINEA = 'D8DCD9'

def localidades(g):
    out = []
    for c in g['candidaturas'].split(' || '):
        if c.startswith('JAL · '):
            loc = c.split('·')[1].strip()
            if loc not in out: out.append(loc)
    return out

def en_dep(g, cod):
    return str(cod) in str(g.get('departamentos', '')).split(' / ')

def deps_nombre(g):
    return ' / '.join(DEPN.get(d, 'dep. ' + d) for d in str(g.get('departamentos', '')).split(' / ') if d)

CASOS = [
    ('Mismo municipio no', lambda g: g['decision'] == 'UNIFICA' and g['tipo'] == 'varios-municipios'),
    ('Misma localidad no', lambda g: g['decision'] == 'UNIFICA' and g['tipo'] == 'mismo-municipio' and len(localidades(g)) > 1),
    ('Varios departamentos', lambda g: g['decision'] != 'UNIFICA'),
]
ETIQUETA = {
    'Mismo municipio no': 'Unificado · municipios distintos',
    'Misma localidad no': 'Unificado · localidades distintas',
    'Varios departamentos': 'Separado · departamentos distintos',
}

def compacta(cand):
    t = cand.replace('PARTIDO ', '').replace('MOVIMIENTO ', '').replace(' COLOMBIANO', '')
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def hoja_instrucciones(wb, depnombre, total, conteos):
    ws = wb.create_sheet('Cómo llenarlo', 0)
    ws.column_dimensions['A'].width = 104
    filas = [
        ('Homónimos de ' + depnombre + ' · base electoral Candidato 360', 16, True, VERDE),
        ('', 11, False, None),
        ('Qué es esto', 12, True, INK),
        ('Nuestra base tiene una fila por CANDIDATURA (439.015 en total). Para mostrar una sola', 11, False, None),
        ('ficha por persona las juntamos cuando el nombre coincide. Con nombres de tres palabras', 11, False, None),
        ('—«Daniel Carvalho Mejía»— eso puede fallar en los dos sentidos: juntar a dos personas', 11, False, None),
        ('distintas, o dejar separada a la misma. Estos son los ' + str(total) + ' casos dudosos de ' + depnombre + '.', 11, False, None),
        ('', 11, False, None),
        ('Qué te pedimos', 12, True, INK),
        ('En la hoja «Casos», columna B, escribe:', 11, False, None),
        ('     X     si TODAS las candidaturas de esa fila son de la MISMA persona', 11, True, VERDE),
        ('     N     si son personas DISTINTAS (homónimos)', 11, True, CORAL),
        ('     ?     si no estás segura (es una respuesta válida y nos sirve)', 11, False, None),
        ('', 11, False, None),
        ('Si en una fila hay tres candidaturas y dos son de la misma persona pero la tercera no,', 11, False, None),
        ('escribe N y cuéntanos en la columna C («Nota») cuál sobra. Con eso arreglamos el caso.', 11, False, None),
        ('', 11, False, None),
        ('Los tres tipos de caso (columna «Caso»)', 12, True, INK),
        ('Unificado · municipios distintos: los juntamos como UNA persona, pero se presentó en', 11, False, None),
        ('     municipios distintos del mismo departamento.  (' + str(conteos[0]) + ' casos)', 11, False, None),
        ('Unificado · localidades distintas: los juntamos como UNA persona, pero fue a JAL de', 11, False, None),
        ('     comunas o localidades distintas.  (' + str(conteos[1]) + ' casos)', 11, False, None),
        ('Separado · departamentos distintos: los dejamos como personas DISTINTAS porque el', 11, False, None),
        ('     nombre aparece en varios departamentos.  (' + str(conteos[2]) + ' casos)', 11, False, None),
        ('', 11, False, None),
        ('Cada fila trae corporación, municipio, año, partido y votos: normalmente el partido y el', 11, False, None),
        ('municipio bastan para saber si es la misma persona. No hace falta llenar todo de una vez,', 11, False, None),
        ('ni investigar: lo que quede en blanco lo revisamos nosotros.', 11, False, None),
    ]
    for i, (texto, tam, negrita, color) in enumerate(filas, start=1):
        c = ws.cell(row=i, column=1, value=texto)
        c.font = Font(name='Calibri', size=tam, bold=negrita, color=color or INK)
        c.alignment = Alignment(vertical='center')
        ws.row_dimensions[i].height = 22 if tam > 12 else 17
    return ws

def recomprime(ruta):
    tmp = ruta + '.tmp'
    with zipfile.ZipFile(ruta) as z, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as o:
        for i in z.infolist(): o.writestr(i.filename, z.read(i.filename))
    shutil.move(tmp, ruta)

def construir(cod, depnombre, salida):
    filas = []; vistos = set(); conteos = []
    for etiqueta, prueba in CASOS:
        antes = len(filas)
        for g in grupos:
            if not en_dep(g, cod) or g['key'] in vistos or not prueba(g): continue
            vistos.add(g['key']); filas.append((etiqueta, g))
        conteos.append(len(filas) - antes)
    wb = Workbook(); wb.remove(wb.active)
    hoja_instrucciones(wb, depnombre, len(filas), conteos)
    ws = wb.create_sheet('Casos')
    cab = ['#', '¿MISMA PERSONA? (X / N / ?)', 'Nota (opcional)', 'Nombre', 'Caso', 'Dónde se presentó', 'Candidaturas', 'Cuántas']
    ws.append(cab)
    for i, c in enumerate(cab, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill('solid', fgColor=INK)
        cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center' if i in (1, 2, 8) else 'left')
    ws.row_dimensions[1].height = 34
    borde = Border(bottom=Side(style='thin', color=LINEA))
    for n, (etiqueta, g) in enumerate(filas, start=1):
        cands = '\n'.join('• ' + compacta(x) for x in g['candidaturas'].split(' || '))
        donde = deps_nombre(g) if etiqueta == 'Varios departamentos' else (g.get('municipios', '') or '—')
        ws.append([n, '', '', g['key'], ETIQUETA[etiqueta], donde, cands, g['n']])
        fila = ws.max_row
        for col in range(1, len(cab) + 1):
            cell = ws.cell(row=fila, column=col)
            cell.border = borde
            cell.alignment = Alignment(vertical='top', wrap_text=col in (3, 5, 6, 7), horizontal='center' if col in (1, 2, 8) else 'left')
            cell.font = Font(size=10, bold=col == 4)
        ws.cell(row=fila, column=2).fill = PatternFill('solid', fgColor='FFF4CC')
        ws.row_dimensions[fila].height = max(16 * g['n'], 30)
    anchos = [5, 17, 26, 30, 26, 26, 76, 9]
    for i, a in enumerate(anchos, start=1): ws.column_dimensions[get_column_letter(i)].width = a
    dv = DataValidation(type='list', formula1='"X,N,?"', allow_blank=True, showDropDown=False)
    dv.prompt = 'X = misma persona · N = personas distintas · ? = no estoy segura'
    dv.promptTitle = '¿Es la misma persona?'
    ws.add_data_validation(dv); dv.add(f'B2:B{ws.max_row}')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:H{ws.max_row}'
    wb.save(salida); recomprime(salida)
    print(os.path.basename(salida), len(filas), 'casos', conteos, os.path.getsize(salida), 'bytes')

def sin_tildes(texto):
    tabla = str.maketrans('áéíóúüñÁÉÍÓÚÜÑ', 'aeiouunAEIOUUN')
    return re.sub(r'[^A-Za-z]', '', texto.translate(tabla))

os.makedirs(SALIDA, exist_ok=True)
for cod in (sys.argv[1:] or ['1']):
    nombre = DEPN.get(cod)
    if not nombre:
        print('no conozco el departamento', cod); continue
    construir(cod, nombre, os.path.join(SALIDA, 'Homonimos_' + sin_tildes(nombre) + '_Candidato360.xlsx'))
