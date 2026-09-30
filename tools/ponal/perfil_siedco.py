#!/usr/bin/env python3
"""Perfila los xlsx por delito de una entrega SIEDCO: fila de encabezado,
columnas, columna de conteo, filas, hechos y rango de fechas."""
import sys, json, re, datetime
from pathlib import Path
from openpyxl import load_workbook

def header_row(rows):
    for i, r in enumerate(rows):
        txt = [c for c in r if isinstance(c, str) and c.strip()]
        if len(txt) >= 8:
            return i
    return None

def count_col(hdr):
    for j, h in enumerate(hdr):
        s = str(h or '').upper()
        if 'CANTIDAD' in s or 'INTERVINIENTES' in s or re.search(r'DEL \d\d/\d\d/\d{4}', s):
            return j
    return None

def perfil(p):
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = []
    for r in it:
        head.append(r)
        if len(head) > 20: break
    h = header_row(head)
    hdr = [str(c).strip() if c is not None else '' for c in head[h]]
    cc = count_col(hdr)
    fcol = next((j for j, x in enumerate(hdr) if 'FECHA' in x.upper()), None)
    rows = head[h+1:] + list(it)
    n = s = 0; fmin = fmax = None; meta = [' | '.join(str(c) for c in r if c) for r in head[:h]]
    for r in rows:
        if r is None or all(c in (None, '') for c in r): continue
        v = r[cc] if cc is not None and cc < len(r) else None
        if not isinstance(v, (int, float)): continue
        n += 1; s += v
        f = r[fcol] if fcol is not None else None
        if isinstance(f, (int, float)): f = datetime.datetime(1899,12,30)+datetime.timedelta(days=f)
        if isinstance(f, datetime.datetime):
            fmin = f if fmin is None or f < fmin else fmin
            fmax = f if fmax is None or f > fmax else fmax
    return dict(archivo=p.name, hdr_row=h+1, cols=hdr, count_col=hdr[cc] if cc is not None else None,
                filas=n, hechos=int(s), fmin=str(fmin)[:10], fmax=str(fmax)[:10], meta=meta[-4:], hojas=len(wb.worksheets))

out = [perfil(p) for p in sorted(Path(sys.argv[1]).glob('*.xlsx')) if not p.name.startswith('~$')]
json.dump(out, open(sys.argv[2], 'w'), ensure_ascii=False, indent=1)
for o in out: print(f"{o['archivo'][:34]:34s} hdr{o['hdr_row']:>3} cols{len(o['cols']):>3} filas{o['filas']:>8,} hechos{o['hechos']:>9,} {o['fmin']}→{o['fmax']} hojas{o['hojas']}")
