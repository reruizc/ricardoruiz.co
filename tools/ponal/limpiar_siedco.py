#!/usr/bin/env python3
"""
limpiar_siedco.py — normaliza una entrega SIEDCO (un .xlsx por delito) a CSV.

Uso:
    python3 tools/ponal/limpiar_siedco.py <dir-xlsx> <periodo> <dir-salida>

    python3 tools/ponal/limpiar_siedco.py "PONAL 2026/2025-v2/DELITOS 2025" 2025 \
        "Bases de datos/PONAL-SIEDCO/clean/2025"

Salida (en <dir-salida>):
    {tipologia}.csv   una fila por registro, columnas fijas (COLS), UTF-8, coma.
                      Mismo esquema que `Bases de datos/Seguridad/Enero 2026/clean/`,
                      que es lo que leen los scripts de seguridad de Medellín.
    calidad.json      por delito: filas, hechos, rango de fechas, columnas que
                      llegaron y % de llenado de cada una.

Las trampas que resuelve (medidas en las entregas de 2025 y 2026):
  · La fila del encabezado NO es fija (11, 12 o 13 según el lote y el archivo)
    → se detecta por contenido: la primera fila con ≥8 celdas de texto.
  · La columna de conteo no se llama igual: `CANTIDAD` en unos lotes y
    `DEL 01/01/2026 AL 31/08/2026` en otros → se renombra a CANTIDAD.
  · Prefijos de cubo (`Hechos.`, `Person.`, …) que cambian por delito → fuera.
  · El nombre del ARCHIVO cambia entre lotes (`HOMICIDIOS AT` / `HOMICIDIOS EN
    AT`, `HURTO A MOTOS` / `HURTO MOTOCICLETAS`…) → mapa `TIPOS`. Un nombre
    desconocido ABORTA: mapearlo mal no falla, deja el delito en cero.
  · Fecha como datetime o serial de Excel; hora como time, datetime o fracción.
  · `CODIGO_DANE` de 8 dígitos (dane5 + '000') → columna `DANE_MUN` de 5.
  · `-` es el nulo de la fuente → celda vacía.
Los microdatos traen edad, sexo, profesión y barrio de víctimas: la salida
es PRIVADA (ver CLAUDE.md, sección PONAL). No subir a un prefijo público.
"""
import csv
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

COLS = [
    "FECHA_HECHO", "ANIO", "MES_NUM", "DIA_SEMANA", "HORA_HECHO",
    "DEPTO_HECHO", "MUNICIPIO_HECHO", "CODIGO_DANE", "DANE_MUN",
    "ZONA", "COMUNAS_ZONAS_DESCRIPCION", "BARRIOS_HECHO", "CLASE_SITIO",
    "ARMAS_MEDIOS", "MOVIL_AGRESOR", "MOVIL_VICTIMA",
    "EDAD", "GENERO", "ESTADO_CIVIL_PERSONA", "PAIS_PERSONA", "PROFESIONES",
    "CARGO_PERSONA", "GRADO_INSTRUCCION_PERSONA", "DESCRIPCION_CONDUCTA",
    "CLASE_BIEN", "TIPO_BIEN", "MARCA_BIEN", "MODELO_VEHICULO", "COLOR_BIEN",
    "CANTIDAD",
]

TIPOS = {
    "AMENAZAS": "amenazas",
    "DELITOS INFORMATICOS": "delitos-informaticos",
    "DELITOS SEXUALES": "delitos-sexuales",
    "EXTORSION": "extorsion",
    "HOMICIDIOS": "homicidios",
    "HOMICIDIOS EN AT": "homicidios-en-at",
    "HOMICIDIOS AT": "homicidios-en-at",
    "HURTO A COMERCIO": "hurto-a-comercio",
    "HURTO A MOTOS": "hurto-a-motos",
    "HURTO MOTOCICLETAS": "hurto-a-motos",
    "HURTO A PERSONAS": "hurto-a-personas",
    "HURTO A RESIDENCIAS": "hurto-a-residencias",
    "HURTO AUTOMOTORES": "hurto-automotores",
    "HURTO BICICLETAS": "hurto-bicicletas",
    "HURTO CABEZAS DE GANADO": "hurto-ganado",
    "HURTO A CABEZAS DE GANADO": "hurto-ganado",
    "HURTO CELULAR": "hurto-celular",
    "HURTO A CELULAR": "hurto-celular",
    "HURTO ENTIDADES FINANCIERAS": "hurto-entidades-financieras",
    "LESIONES EN AT": "lesiones-en-at",
    "LESIONES EN ACCIDENTES DE TRANSITO": "lesiones-en-at",
    "LESIONES PERSONALES": "lesiones-personales",
    "PIRATERIA TERRESTRE": "pirateria-terrestre",
    "HURTO PIRATERIA TERRESTRE": "pirateria-terrestre",
    "SECUESTRO": "secuestro",
    "TERRORISMO": "terrorismo",
    "VIOLENCIA INTRAFAMILIAR": "violencia-intrafamiliar",
}

RE_PREF = re.compile(r"^(Hechos|Person|Calend|Conduc|ClaseS|Bienes|ListaC|Armas|Hechos1)\.", re.I)
NULOS = {"", "-", "--", "NULL", "N/A"}
EXCEL0 = dt.datetime(1899, 12, 30)


def sin_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def tipo_de(archivo):
    k = re.sub(r"\s+", " ", sin_tildes(Path(archivo).stem).upper()).strip()
    if k not in TIPOS:
        sys.exit(f"nombre de archivo desconocido: {archivo!r} — agregarlo a TIPOS")
    return TIPOS[k]


def col_canon(h):
    s = RE_PREF.sub("", str(h or "").strip())
    u = sin_tildes(s).upper()
    if "CANTIDAD" in u or "INTERVINIENTES" in u or re.search(r"DEL \d\d/\d\d/\d{4}", u):
        return "CANTIDAD"
    if u == "DEPARTAMENTO":
        return "DEPTO_HECHO"
    return u


def a_fecha(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return EXCEL0 + dt.timedelta(days=float(v))
    if isinstance(v, str):
        for f in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return dt.datetime.strptime(v.strip()[:19], f)
            except ValueError:
                pass
    return None


def a_hora(v):
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M:%S")
    if isinstance(v, (int, float)) and 0 <= v < 1:
        s = round(v * 86400)
        return f"{s // 3600 % 24:02d}:{s // 60 % 60:02d}:{s % 60:02d}"
    if isinstance(v, str) and re.match(r"^\d{1,2}:\d\d", v.strip()):
        p = v.strip().split(":")
        return f"{int(p[0]):02d}:{p[1][:2]}:{(p[2][:2] if len(p) > 2 else '00')}"
    return ""


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s.upper() in NULOS else s


def limpiar(path, periodo, out_dir):
    tipo = tipo_de(path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = None
    for r in it:
        if len([c for c in r if isinstance(c, str) and c.strip()]) >= 8:
            hdr = [col_canon(c) for c in r]
            break
    if hdr is None or "CANTIDAD" not in hdr:
        sys.exit(f"{path.name}: no encontré encabezado con columna de conteo")
    ix = {}
    for j, c in enumerate(hdr):
        if c and c not in ix:
            ix[c] = j
    llegaron = sorted(c for c in ix if c)
    fuera = [c for c in llegaron if c not in COLS and c not in ("MES",)]

    llenas = {c: 0 for c in COLS}
    filas = hechos = 0
    fmin = fmax = None
    fuera_periodo = 0
    out = out_dir / f"{tipo}.csv"
    with out.open("w", encoding="utf-8", newline="") as fo:
        w = csv.writer(fo)
        w.writerow(COLS)
        for r in it:
            if r is None:
                continue
            g = lambda c: r[ix[c]] if c in ix and ix[c] < len(r) else None
            q = g("CANTIDAD")
            if not isinstance(q, (int, float)):
                continue            # filas de total / pie de tabla
            q = int(q) or 1
            f = a_fecha(g("FECHA_HECHO"))
            if f is not None:
                fmin = f if fmin is None or f < fmin else fmin
                fmax = f if fmax is None or f > fmax else fmax
                if str(f.year) != str(periodo)[:4]:
                    fuera_periodo += q
            dane = re.sub(r"\D", "", texto(g("CODIGO_DANE")))
            fila = {c: texto(g(c)) for c in COLS if c in ix}
            fila.update({
                "FECHA_HECHO": f.strftime("%Y-%m-%d") if f else "",
                "ANIO": f.year if f else "",
                "MES_NUM": f.month if f else "",
                "HORA_HECHO": a_hora(g("HORA_HECHO")),
                "CODIGO_DANE": dane,
                "DANE_MUN": dane[:-3].zfill(5) if len(dane) >= 7 else "",
                "CANTIDAD": q,
            })
            for c in COLS:
                if fila.get(c) not in (None, ""):
                    llenas[c] += 1
            w.writerow([fila.get(c, "") for c in COLS])
            filas += 1
            hechos += q
    return tipo, {
        "archivo": path.name, "filas": filas, "hechos": hechos,
        "fecha_min": fmin.strftime("%Y-%m-%d") if fmin else None,
        "fecha_max": fmax.strftime("%Y-%m-%d") if fmax else None,
        "hechos_fuera_del_periodo": fuera_periodo,
        "columnas_recibidas": llegaron,
        "columnas_no_previstas": fuera,
        "llenado_pct": {c: round(100 * n / filas, 1) for c, n in llenas.items() if filas and c in ix},
        "ausentes": [c for c in COLS if c not in ix and c not in
                     ("ANIO", "MES_NUM", "DANE_MUN")],
    }


def rellenar_dane(out_dir, cal):
    """Terrorismo llega SIN `CODIGO_DANE` (solo depto + municipio en texto).
    Se reconstruye con la moda del par (depto, municipio) → DANE en los demás
    archivos de la MISMA entrega. Un par que no aparezca en ningún otro delito
    queda vacío: no se adivina por parecido (`El Carmen` existe en Norte de
    Santander y en Bolívar)."""
    from collections import Counter, defaultdict
    moda = defaultdict(Counter)
    faltan = [t for t, d in cal.items() if "CODIGO_DANE" in d["ausentes"]]
    if not faltan:
        return
    for p in out_dir.glob("*.csv"):
        if p.stem in faltan:
            continue
        with p.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r["DANE_MUN"]:
                    moda[(r["DEPTO_HECHO"].upper(), r["MUNICIPIO_HECHO"].upper())][r["CODIGO_DANE"]] += 1
    for t in faltan:
        p = out_dir / f"{t}.csv"
        rows = list(csv.DictReader(p.open(encoding="utf-8")))
        ok = 0
        for r in rows:
            c = moda.get((r["DEPTO_HECHO"].upper(), r["MUNICIPIO_HECHO"].upper()))
            if c:
                r["CODIGO_DANE"] = c.most_common(1)[0][0]
                r["DANE_MUN"] = r["CODIGO_DANE"][:-3].zfill(5)
                ok += 1
        with p.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=COLS)
            w.writeheader()
            w.writerows(rows)
        cal[t]["dane_reconstruido"] = f"{ok} de {len(rows)} filas, por (depto, municipio) de los demás delitos"
        print(f"  {t}: DANE reconstruido en {ok}/{len(rows)} filas")


def main():
    src, periodo, out_dir = Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3])
    out_dir.mkdir(parents=True, exist_ok=True)
    cal = {}
    for p in sorted(src.glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        t0 = dt.datetime.now()
        tipo, info = limpiar(p, periodo, out_dir)
        if tipo in cal:
            sys.exit(f"{tipo} aparece dos veces en la entrega")
        cal[tipo] = info
        print(f"  {tipo:28s} {info['filas']:>8,} filas {info['hechos']:>8,} hechos "
              f"{info['fecha_min']}→{info['fecha_max']}  ({(dt.datetime.now()-t0).seconds}s)")
    rellenar_dane(out_dir, cal)
    (out_dir / "calidad.json").write_text(json.dumps(
        {"periodo": periodo, "fuente": str(src), "generado": dt.datetime.now().isoformat(timespec="seconds"),
         "delitos": cal}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"  total {sum(v['hechos'] for v in cal.values()):,} hechos en {len(cal)} delitos")


if __name__ == "__main__":
    main()
