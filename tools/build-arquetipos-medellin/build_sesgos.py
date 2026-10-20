#!/usr/bin/env python3
"""
build-arquetipos-medellin/build_sesgos.py

Sesgos y palancas de los 5 arquetipos 2027 de Medellín (tablero de Nury,
entrega del 1-oct-2026) → sesgos-palancas.json, que lee
proyecto-dc/arquetipos.html.

Uso:
    python3 tools/build-arquetipos-medellin/build_sesgos.py \\
        "Bases de datos/proyecto-dc/entrega-nury-medellin-20261001/ACTUALIZADOS/TABLERO_ARQUETIPOS_MEDELLIN_2027_SESGOS_PALANCAS_FINAL.xlsx" \\
        "Bases de datos/proyecto-dc/arquetipos"

Del tablero se leen tres hojas:
  · 00_Dashboard  → la sensibilidad 1-5 de cada palanca (la vigente).
  · 01_Sesgos     → 5 sesgos por arquetipo. ⚠️ El encabezado está corrido:
                    dice «Tipo · Qué significa · Cómo puede aparecer», pero las
                    columnas traen qué significa, cómo aparece y cómo analizarlo.
  · 02_Palancas_11 → qué pasa si cada palanca empeora o mejora, y la posible
                    transición del arquetipo. ⚠️ Su columna de sensibilidad es
                    la VIEJA: la cifra sale del dashboard, no de aquí.

⚠️ La palanca 11 trae el nombre del candidato en el rótulo. Se publica como
«Huella territorial del equipo» y el script aborta si el nombre sobrevive.

Solo diagnóstico agregado: el propio tablero advierte que los valores no son
intención de voto ni eficacia de un mensaje.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

# El nombre del arquetipo 2027 → id de familia de arquetipos.json.
FAMILIA = {
    'Resultados Primero': 'proteccion',
    'Ojo Abierto': 'castigo',
    'Confianza Cercana': 'continuidad',
    'Hogar Seguro': 'supervivencia',
    'Orgullo de Barrio': 'pertenencia',
}
PALANCA_PUBLICA = {re.compile(r'^Huella territorial del equipo\b.*', re.I): 'Huella territorial del equipo'}
# Los términos que delatan el caso viven FUERA del repo (es público), uno por
# línea. Sin la lista el script no corre: falla cerrado.
PROHIBIDO_TXT = 'Bases de datos/proyecto-dc/prohibido.txt'


def limpio(s):
    s = re.sub(r'\s+', ' ', str(s or '')).strip()
    for pat, rep in PALANCA_PUBLICA.items():
        if pat.match(s):
            return rep
    return s


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    xlsx, out = Path(sys.argv[1]), Path(sys.argv[2])
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # 1. Sensibilidad vigente: la tabla del dashboard (palanca × arquetipo).
    d = list(wb['00_Dashboard'].iter_rows(values_only=True))
    cab = next(i for i, r in enumerate(d) if r[0] == 'Palanca')
    nombres = [limpio(x) for x in d[cab][1:] if x]
    if set(nombres) != set(FAMILIA):
        sys.exit(f'arquetipos del dashboard inesperados: {nombres}')
    palancas, sens = [], {}
    for r in d[cab + 1:]:
        if not r[0]:
            break
        p = limpio(r[0])
        palancas.append(p)
        for n, v in zip(nombres, r[1:1 + len(nombres)]):
            sens.setdefault(FAMILIA[n], {})[p] = int(v)
    if len(palancas) != 11:
        sys.exit(f'se esperaban 11 palancas y hay {len(palancas)}')

    # 2. Los sesgos.
    sesgos = {f: [] for f in FAMILIA.values()}
    for r in wb['01_Sesgos'].iter_rows(min_row=5, values_only=True):
        n = limpio(r[0])
        if n not in FAMILIA or not r[1]:
            continue
        sesgos[FAMILIA[n]].append({'sesgo': limpio(r[1]), 'significa': limpio(r[2]), 'aparece': limpio(r[3]), 'analizar': limpio(r[4])})

    # 3. Si empeora / si mejora, y la transición.
    reaccion, transicion = {f: {} for f in FAMILIA.values()}, {}
    for r in wb['02_Palancas_11'].iter_rows(min_row=5, values_only=True):
        n = limpio(r[0])
        if n not in FAMILIA or not r[1]:
            continue
        f, p = FAMILIA[n], limpio(r[1])
        if p not in palancas:
            sys.exit(f'palanca desconocida en 02_Palancas_11: {p}')
        reaccion[f][p] = {'empeora': limpio(r[3]), 'mejora': limpio(r[4])}
        transicion[f] = limpio(r[5])

    for f in FAMILIA.values():
        if len(sesgos[f]) != 5:
            sys.exit(f'{f}: {len(sesgos[f])} sesgos (se esperaban 5)')
        if set(reaccion[f]) != set(palancas):
            sys.exit(f'{f}: faltan palancas en 02_Palancas_11')

    datos = {
        'v': '2026-10-20',
        'fuente': 'Tablero psicopolítico de arquetipos Medellín 2027 (entrega del 1-oct-2026).',
        'escala': 'Sensibilidad 1-5: qué tan probable es que un cambio en la palanca altere la atención o la evaluación política del arquetipo. No es intención de voto ni eficacia de un mensaje.',
        'palancas': palancas,
        'arquetipos': {f: {'nombre': n, 'sensibilidad': sens[f], 'sesgos': sesgos[f], 'reaccion': reaccion[f], 'transicion': transicion[f]}
                       for n, f in FAMILIA.items()},
    }
    txt = json.dumps(datos, ensure_ascii=False, indent=1)
    lista = Path(PROHIBIDO_TXT)
    if not lista.exists():
        sys.exit(f'falta {PROHIBIDO_TXT} (los términos del caso que no pueden publicarse)')
    terminos = [t.strip() for t in lista.read_text(encoding='utf-8').splitlines() if t.strip()]
    if any(re.search(re.escape(t), txt, re.I) for t in terminos):
        sys.exit('⚠️ el JSON trae el nombre del candidato: no se escribe')
    out.mkdir(parents=True, exist_ok=True)
    (out / 'sesgos-palancas.json').write_text(txt + '\n', encoding='utf-8')
    print(f'→ {out / "sesgos-palancas.json"} · {len(palancas)} palancas · 25 sesgos · {len(txt) // 1024} KB')


if __name__ == '__main__':
    main()
