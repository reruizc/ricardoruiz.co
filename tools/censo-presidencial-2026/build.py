#!/usr/bin/env python3
"""Censo electoral PRESIDENCIAL 2026 por puesto, con sexo y 10 rangos de edad → Excel.

Fuente: el JSON de "Consulta los censos históricos" de la Registraduría
(https://www.registraduria.gov.co/IMG/json/2026_datos_censo-electoral-min.json),
bajado el 26-sep-2026 a Bases de datos/censos-registraduria/. Solo se lee desde
el navegador (reto de Cloudflare): ver CLAUDE.md, sección del escrutinio en CSV.

Salida: Bases de datos/output_censo_2026/Censo_Presidencial_2026_por_puesto.xlsx
Uso:    python3 tools/censo-presidencial-2026/build.py
"""
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parents[2]
FUENTE = RAIZ / 'Bases de datos/censos-registraduria/2026_datos_censo-electoral-min.json'
SALIDA = RAIZ / 'Bases de datos/output_censo_2026/Censo_Presidencial_2026_por_puesto.xlsx'

# Controles: si la fuente cambia, el Excel no se arma en silencio con otra cifra.
TOTAL_PRES = 41_421_973
TOTAL_CONG = 41_287_084

# La fuente recorta dos nombres de departamento.
DEP_NOMBRE = {'25': 'NORTE DE SANTANDER', '31': 'VALLE DEL CAUCA', '88': 'CONSULADOS (EXTERIOR)'}

EDADES = ['18-20', '21-25', '26-30', '31-35', '36-40', '41-45', '46-50',
          '51-55', '56-60', 'mas de 60', 'sin fecha de nacimiento']
ROT_EDAD = {'mas de 60': 'Más de 60', 'sin fecha de nacimiento': 'Sin fecha de nac.'}

TINTA = '1A1A2E'
CABECERA = PatternFill('solid', fgColor=TINTA)
BLANCO = Font(bold=True, color='FFFFFF')
MILES = '#,##0'
PCT = '0.0%'


def n(x, signo=False):
    """Número con punto de miles (convención colombiana)."""
    return (f'{x:+,}' if signo else f'{x:,}').replace(',', '.')


def cargar(tipo):
    filas = json.load(open(FUENTE, encoding='utf-8'))['data']
    return [r for r in filas if r['tipo_eleccion'].startswith(tipo)]


def dep_nombre(r):
    return DEP_NOMBRE.get(r['cod_depto'], r['departamento'])


def fila_valores(pot, muj, hom, edades):
    return [pot, muj, hom, (muj / pot if pot else 0)] + [edades.get(e, 0) for e in EDADES]


def hoja_tabla(wb, titulo, nombre_tabla, cabeceras, filas, cols_texto):
    ws = wb.create_sheet(titulo)
    ws.append(cabeceras)
    for f in filas:
        ws.append(f)
    for c in ws[1]:
        c.font = BLANCO
        c.fill = CABECERA
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 32
    for i, cab in enumerate(cabeceras, start=1):
        letra = get_column_letter(i)
        if i <= cols_texto:
            ancho = max(len(str(cab)), *(len(str(f[i - 1])) for f in filas[:3000])) + 2
            ws.column_dimensions[letra].width = min(ancho, 48)
        else:
            ws.column_dimensions[letra].width = 12
            fmt = PCT if cab == '% mujeres' else MILES
            for celda in ws[letra][1:]:
                celda.number_format = fmt
    ref = f'A1:{get_column_letter(len(cabeceras))}{len(filas) + 1}'
    tabla = Table(displayName=nombre_tabla, ref=ref)
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = ws.cell(row=2, column=cols_texto + 1)
    return ws


def main():
    pres = cargar('PRESIDENTE')
    cong = cargar('CONGRESO')

    total = sum(r['potencial_electoral'] for r in pres)
    assert total == TOTAL_PRES, f'El censo presidencial suma {total:,}, se esperaba {TOTAL_PRES:,}'
    assert sum(r['potencial_electoral'] for r in cong) == TOTAL_CONG
    claves = {(r['cod_depto'], r['cod_mun'], r['zona'], r['cod_puesto']) for r in pres}
    assert len(claves) == len(pres), 'Hay puestos duplicados'
    for r in pres:
        assert sum(r['sexo'].values()) == r['potencial_electoral'], r
        assert sum(r['rango_etario'].values()) == r['potencial_electoral'], r

    pres.sort(key=lambda r: (r['cod_depto'], r['cod_mun'], r['zona'], r['cod_puesto']))
    cab_valores = ['Potencial electoral', 'Mujeres', 'Hombres', '% mujeres'] + \
                  [ROT_EDAD.get(e, e) for e in EDADES]

    # ── por puesto ──
    filas_p = []
    for r in pres:
        s = r['sexo']
        filas_p.append([
            r['cod_depto'] + r['cod_mun'] + r['zona'] + r['cod_puesto'],
            r['cod_depto'], dep_nombre(r), r['cod_mun'], r['municipio'],
            r['zona'], r['cod_puesto'], r['puesto_votacion'],
        ] + fila_valores(r['potencial_electoral'], s['mujeres'], s['hombres'], r['rango_etario']))

    # ── agregados por municipio y por departamento ──
    def agregar(llave, rotulo):
        acc = defaultdict(lambda: {'rot': None, 'pue': 0, 'pot': 0, 'm': 0, 'h': 0,
                                   'e': defaultdict(int)})
        for r in pres:
            a = acc[llave(r)]
            a['rot'] = rotulo(r)
            a['pue'] += 1
            a['pot'] += r['potencial_electoral']
            a['m'] += r['sexo']['mujeres']
            a['h'] += r['sexo']['hombres']
            for k, v in r['rango_etario'].items():
                a['e'][k] += v
        return [a['rot'] + [a['pue']] + fila_valores(a['pot'], a['m'], a['h'], a['e'])
                for _, a in sorted(acc.items())]

    filas_m = agregar(lambda r: (r['cod_depto'], r['cod_mun']),
                      lambda r: [r['cod_depto'] + r['cod_mun'], r['cod_depto'], dep_nombre(r),
                                 r['cod_mun'], r['municipio']])
    filas_d = agregar(lambda r: r['cod_depto'],
                      lambda r: [r['cod_depto'], dep_nombre(r)])

    # ── Excel ──
    wb = Workbook()
    lee = wb.active
    lee.title = 'Léeme'
    ext_p = sum(r['potencial_electoral'] for r in pres if r['cod_depto'] == '88')
    ext_c = sum(r['potencial_electoral'] for r in cong if r['cod_depto'] == '88')
    texto = [
        ('Censo electoral · Presidencial 2026', Font(bold=True, size=16)),
        ('Potencial electoral por puesto de votación, con sexo y rangos de edad', Font(italic=True, size=11)),
        ('', None),
        ('Qué es', Font(bold=True, size=12)),
        ('El censo con el que se votó la elección presidencial de 2026 (primera y segunda vuelta), '
         'tal como lo publica la Registraduría Nacional del Estado Civil en su consulta de censos '
         'históricos. Es el censo electoral más reciente de alcance nacional.', None),
        ('', None),
        ('Cifras', Font(bold=True, size=12)),
        (f'Potencial electoral total: {n(TOTAL_PRES)}', None),
        (f'En Colombia: {n(TOTAL_PRES - ext_p)} · En el exterior: {n(ext_p)}', None),
        (f'Puestos de votación: {n(len(pres))} (incluidos {sum(1 for r in pres if r["cod_depto"] == "88")} en el exterior)', None),
        ('', None),
        ('Frente al censo del Congreso (marzo de 2026)', Font(bold=True, size=12)),
        (f'El censo del Congreso era de {n(TOTAL_CONG)}. El presidencial tiene '
         f'{n(TOTAL_PRES - TOTAL_CONG)} más, pero toda la diferencia está en el exterior '
         f'({n(ext_p - ext_c, signo=True)}). Dentro del país el censo presidencial es '
         f'{n((TOTAL_CONG - ext_c) - (TOTAL_PRES - ext_p))} menor.', None),
        ('', None),
        ('Hojas', Font(bold=True, size=12)),
        ('Por puesto: una fila por puesto de votación. Código de puesto = depto (2) + municipio (3) + zona (2) + puesto (2).', None),
        ('Por municipio y Por departamento: los mismos datos sumados.', None),
        ('', None),
        ('Tenga en cuenta', Font(bold=True, size=12)),
        ('· Los códigos son de la Registraduría, no del DANE (Antioquia = 01, Bogotá = 16).', None),
        ('· El código de puesto puede llevar letras (p. ej. A1 en el exterior y en algunos puestos de Nariño y Bogotá): trátelo como texto.', None),
        ('· No trae número de mesas ni direcciones; para eso use el Divipole del Congreso de marzo de 2026.', None),
        ('· La edad es la que registra la Registraduría al cierre del censo. "Sin fecha de nacimiento" son registros sin ese dato.', None),
        ('· La fuente no registra votantes en la categoría "otros" de sexo; por eso no hay columna.', None),
        ('', None),
        ('Fuente: Registraduría Nacional del Estado Civil · consulta de censos históricos '
         '(registraduria.gov.co). Procesado por ricardoruiz.co.', Font(italic=True, size=9)),
    ]
    for i, (t, f) in enumerate(texto, start=1):
        c = lee.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if f:
            c.font = f
    lee.column_dimensions['A'].width = 110

    hoja_tabla(wb, 'Por puesto', 'CensoPuesto',
               ['Código puesto', 'Cód. depto', 'Departamento', 'Cód. municipio', 'Municipio',
                'Zona', 'Puesto', 'Nombre del puesto'] + cab_valores,
               filas_p, cols_texto=8)
    hoja_tabla(wb, 'Por municipio', 'CensoMunicipio',
               ['Código', 'Cód. depto', 'Departamento', 'Cód. municipio', 'Municipio', 'Puestos']
               + cab_valores, filas_m, cols_texto=5)
    hoja_tabla(wb, 'Por departamento', 'CensoDepartamento',
               ['Cód. depto', 'Departamento', 'Puestos'] + cab_valores, filas_d, cols_texto=2)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(f'{SALIDA}  ·  {len(filas_p)} puestos · {len(filas_m)} municipios · '
          f'{len(filas_d)} departamentos · {SALIDA.stat().st_size / 1e6:.1f} MB')


if __name__ == '__main__':
    main()
